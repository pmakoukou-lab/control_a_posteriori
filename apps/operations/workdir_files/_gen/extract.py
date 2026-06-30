# -*- coding: utf-8 -*-
"""Outil de DEV (non embarqué) : extrait la structure d'une grille de saisie +
son fichier « base de données » (op/lot) pour générer/vérifier le modèle d'un mode.

Usage :
    python extract.py align <mode>      # vérifie l'alignement grille<->base
    python extract.py gen   <mode>      # génère models/audits/model_<mode>.py
"""
import os, re, sys, unicodedata
import openpyxl
from openpyxl.utils import get_column_letter

GEN_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.dirname(os.path.dirname(GEN_DIR))          # .../operations
GRILLES = os.path.join(APP_DIR, 'grilles')
WORKDIR = os.path.join(APP_DIR, 'workdir_files')
MODELS_OUT = os.path.join(APP_DIR, 'models', 'audits')

# Déclaration des modes : grille de saisie + base de données (op/lot) + export.
MODES = {
    'aor_pi': dict(
        grille=os.path.join(GRILLES, '2- Saisie_AOR nv prestations intellectuelles.xlsx'),
        grille_sheet='Nouvelle grille',
        base=os.path.join(WORKDIR, 'base de données -AOR de prestations intellectuelles_vf.xlsx'),
        label="Appel d'offres restreint — prestations intellectuelles",
        pill='bg-sky-100 text-sky-700',
        export=dict(file='2- Saisie_AOR nv prestations intellectuelles.xlsx', sheet='Nouvelle grille', start_row=7),
    ),
    'aor_ft': dict(
        grille=os.path.join(GRILLES, '2- saisie_AOR_nv fournitures & travaux.xlsx'),
        grille_sheet='Nouvelle grille',
        base=os.path.join(WORKDIR, 'base de données -AOR de fournitures et travaux_vf.xlsx'),
        label="Appel d'offres restreint — fournitures & travaux",
        pill='bg-teal-100 text-teal-700',
        export=dict(file='2- saisie_AOR_nv fournitures & travaux.xlsx', sheet='Nouvelle grille', start_row=7),
    ),
    'gag': dict(
        grille=os.path.join(GRILLES, '3-saisie_GAG_nv.xlsx'),
        grille_sheet='Nouvelle grille',
        base=os.path.join(WORKDIR, 'base de données -GAG_vf.xlsx'),
        label="Gré à gré",
        pill='bg-amber-100 text-amber-700',
        export=dict(file='3-saisie_GAG_nv.xlsx', sheet='Nouvelle grille', start_row=7),
    ),
}

STOP = {'de', 'du', 'des', 'la', 'le', 'les', "l", 'd', 'a', 'au', 'aux', 'et',
        'par', 'pour', 'dans', 'en', 'un', 'une', 'sur', 'ou', 'of', 'the'}


def norm(s):
    if s is None:
        return ''
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()


def slug(s, maxwords=5):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().lower()
    toks = [t for t in re.split(r'[^a-z0-9]+', s) if t and t not in STOP]
    return '_'.join(toks[:maxwords]) or 'champ'


def keynorm(s):
    """Clé de lookup CANON : ascii, minuscule, ponctuation (apostrophes, slashes,
    parenthèses…) réduite à des espaces. Rend la correspondance robuste."""
    return re.sub(r'[^a-z0-9]+', ' ', norm(s)).strip()


# --- Champs « canoniques » : noms attendus par l'engine (collecte / UI lots). ---
# Clé = (keynorm(section), keynorm(sous-section)). '' = pas de sous-section.
CANON = {
    ('controleur', 'nom de l agent dgmp'): ('agent_dgmp', 'string'),
    ('n controleur', 'nom de l agent dgmp'): ('agent_dgmp', 'string'),
    ('date de controle', ''): ('date_controle', 'date'),
    ('ministere', ''): ('ministere', 'reference ministere'),
    ('autorite contractante', ''): ('autorite_contractante', 'reference autorite_contractante'),
    ('personne ressource', 'interlocuteur chez l ac'): ('interlocuteur_ac', 'string'),
    ('imputation budgetaire', ''): ('imputation_budgetaire', 'string'),
    ('montant de la dotation du credit', 'montant'): ('montant_dotation_credit', 'double'),
    ('numero de l operation', 'n de l aor'): ('n_bomp', 'string'),
    ('numero de l operation', 'n de l aoo'): ('n_bomp', 'string'),
    ('numero de l operation', ''): ('n_bomp', 'string'),
    ('objet de l operation', ''): ('objet_operation', 'string'),
    ('objet de l operation', 'observations'): ('objet_operation_observation', 'text'),
    ('nature de l operation', 'travaux fournitures prestations'): ('nature_operation', 'reference nature_operation'),
    ('nature de l operation', 'observations'): ('nature_operation_observation', 'text'),
    # niveau lot (allotissement)
    ('montant de l estimation', 'montant'): ('montant_estimation', 'double'),
    ('attributaire s', 'raison sociale'): ('attributaire_nom', 'string'),
    ('attributaire s', 'ncc attributaire'): ('attributaire_ncc', 'string'),
    ('attributaire s', 'lot'): ('attributaire_lot', 'string'),
    ('attributaire s', 'objet'): ('attributaire_objet', 'string'),
    ('attributaire s', 'montant'): ('attributaire_montant', 'double'),
    ('attributaire s', 'observations'): ('attributaire_observation', 'text'),
    ('montant du marche', 'montant'): ('montant_marche', 'double'),
    ('ecart entre montant estimatif et montant attribue', 'montant'): ('ecart_entre_montant_estimatif', 'double'),
}


def infer_type(section, sub):
    s = norm(sub) or norm(section)
    sec = norm(section)
    if 'reference nature' in sec:
        return 'reference nature_operation'
    if re.search(r'\bdate\b', s):
        return 'date'
    if 'montant' in s or 'cout' in s:
        return 'double'
    if s in ('observations', 'observation', 'nb') or 'observation' in s:
        return 'text'
    if any(k in s for k in ('constat', 'existence', 'respect', 'conformite', 'conforme',
                            'exhaustivite', 'presence')):
        return 'boolean'
    if 'nombre' in s or 'duree' in s:
        return 'integer'
    return 'string'


def grille_blocks(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    blocks, sec, cur = [], None, None
    for ci in range(1, ws.max_column + 1):
        s_raw = ws.cell(row=5, column=ci).value          # en-tête de section (fusionnée)
        sub = ws.cell(row=6, column=ci).value             # sous-section (par colonne)
        if s_raw not in (None, ''):
            sec = str(s_raw)
        # Colonne hors tableau : ni en-tête de section À CETTE colonne, ni sous-section.
        if (s_raw in (None, '')) and (sub in (None, '')):
            continue
        if sec is None:
            continue
        if cur is None or norm(sec) != norm(cur['sec']):
            cur = {'sec': sec, 'cols': []}
            blocks.append(cur)
        cur['cols'].append((get_column_letter(ci), sub))
    wb.close()
    return blocks


def _toks(s):
    return {t for t in re.split(r'[^a-z0-9]+', norm(s)) if t and t not in STOP}


def match_lot_flag(section, base):
    """Associe une section de grille à la section de base la plus proche
    (similarité de Jaccard sur les tokens) et renvoie (is_lot, base_name, score)."""
    gt = _toks(section)
    best, bestname, bestlot = 0.0, '', False
    for bname, is_lot in base:
        bt = _toks(bname)
        if not gt or not bt:
            continue
        j = len(gt & bt) / len(gt | bt)
        if j > best:
            best, bestname, bestlot = j, bname, is_lot
    return bestlot, bestname, best


def base_sections(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    hr = lotc = None
    for r in range(1, min(6, ws.max_row) + 1):
        rowvals = [norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)]
        if 'operation' in rowvals and 'lots' in rowvals:
            hr = r
            lotc = rowvals.index('lots') + 1
            break
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        nm = ws.cell(row=r, column=2).value
        if nm in (None, ''):
            continue
        out.append((str(nm), norm(ws.cell(row=r, column=lotc).value) == 'oui'))
    wb.close()
    return out


def _trim_blocks(g):
    """Retire le bloc d'en-tête « N° » (numéro d'ordre, non stocké)."""
    if g and norm(g[0]['sec']) in ('n', 'no', 'n0', ''):
        return g[1:]
    return g


def align(mode):
    cfg = MODES[mode]
    g = _trim_blocks(grille_blocks(cfg['grille'], cfg['grille_sheet']))
    b = base_sections(cfg['base'])
    print('mode=%s  grille blocks=%d  base sections=%d' % (mode, len(g), len(b)))
    for i in range(max(len(g), len(b))):
        gs = norm(g[i]['sec'])[:30] if i < len(g) else '---'
        bs = (norm(b[i][0])[:26] + ('  LOT' if b[i][1] else '')) if i < len(b) else '---'
        mark = '' if gs[:10] == norm(b[i][0])[:10] else '   <<DIFF' if i < len(b) and i < len(g) else ''
        print('%2d  %-32s | %-32s%s' % (i, gs, bs, mark))


def classify(mode, threshold=0.34):
    """[(block, is_lot, base_name, score), …] : chaque section de grille reliée à
    sa section de base (op/lot) par similarité de tokens. Sous le seuil → op."""
    cfg = MODES[mode]
    g = _trim_blocks(grille_blocks(cfg['grille'], cfg['grille_sheet']))
    b = base_sections(cfg['base'])
    out = []
    for blk in g:
        is_lot, bname, score = match_lot_flag(blk['sec'], b)
        if score < threshold:
            is_lot, bname = False, '(défaut op, score=%.2f)' % score
        out.append((blk, is_lot, bname, score))
    return out


def match(mode):
    for blk, is_lot, bname, score in classify(mode):
        print('%-34s -> %-30s %s  %.2f' % (
            norm(blk['sec'])[:34], norm(bname)[:30], 'LOT ' if is_lot else 'op  ', score))


def build_fields(mode):
    """Retourne (op_fields, lot_fields) : listes de dicts {code,type,label,section,col}."""
    op_fields, lot_fields = [], []
    used = {'op': set(), 'lot': set()}
    for blk, is_lot, _bname, _score in classify(mode):
        bucket = 'lot' if is_lot else 'op'
        for col, sub in blk['cols']:
            section = blk['sec']
            key = (keynorm(section), keynorm(sub))
            if key in CANON:
                code, typ = CANON[key]
            else:
                base_code = slug(section)
                sub_slug = slug(sub) if sub not in (None, '') and norm(sub) != norm(section) else ''
                code = (base_code + ('_' + sub_slug if sub_slug else ''))[:54]
                typ = infer_type(section, sub)
            # unicité par table
            seen = used[bucket]
            c, n = code, 2
            while c in seen:
                c = '%s_%d' % (code, n)
                n += 1
            seen.add(c)
            label = (str(sub).strip() if sub not in (None, '') else str(section).strip())
            label = re.sub(r'\s+', ' ', label)
            (lot_fields if is_lot else op_fields).append(
                dict(code=c, type=typ, label=label,
                     section=re.sub(r'\s+', ' ', str(section).strip()), col=col))
    return op_fields, lot_fields


def _field_line(f):
    label = f['label'].replace('"', '\\"')
    section = f['section'].replace('"', '\\"')
    return ('        Field("%s", "%s", label="%s", comment="Section: %s | col. %s"),'
            % (f['code'], f['type'], label, section, f['col']))


def gen(mode):
    cfg = MODES[mode]
    op_fields, lot_fields = build_fields(mode)
    op_lines = '\n'.join(_field_line(f) for f in op_fields)
    lot_lines = '\n'.join(_field_line(f) for f in lot_fields)
    has_attr = any(f['code'].startswith('attributaire_') for f in lot_fields)
    src = TEMPLATE.format(
        mode=mode, MODE=mode.upper(), label=cfg['label'],
        optbl=mode + '_operations', lottbl=mode + '_allotissements',
        op_lines=op_lines, lot_lines=lot_lines,
        n_op=len(op_fields), n_lot=len(lot_fields),
    )
    out = os.path.join(MODELS_OUT, 'model_%s.py' % mode)
    with open(out, 'w', encoding='utf-8') as fh:
        fh.write(src)
    print('written', out, '| op=%d lot=%d' % (len(op_fields), len(lot_fields)))


TEMPLATE = '''# -*- coding: utf-8 -*-
"""
Modèle pyDAL — Mode de passation : {MODE}.

Fichier généré depuis la grille de saisie + le fichier « base de données »
(workdir_files) qui indique, par section, si elle relève de l'opération ou de
l'allotissement (lot). Tables PROPRES au mode (préfixe ``{mode}_``) :

  - ``{optbl}``     : sections au niveau de l'opération.
  - ``{lottbl}`` : sections renseignées PAR LOT (relation 1-N).

Les tables de référence (ministere, autorite_contractante, nature_operation…)
sont définies en amont par ``table_parametres_models.define_all``.
"""
from datetime import datetime, timezone

from pydal import Field

from ...common import auth


def _meta_fields():
    """Champs de traçabilité communs (soft delete + audit)."""
    return [
        Field("is_deleted", "boolean", default=False, readable=False),
        Field("created_on", "datetime", default=lambda: datetime.now(timezone.utc),
              writable=False),
        Field("created_by", "reference auth_user", default=auth.user_id, writable=False),
    ]


def define_operations(db):
    """Définit la table `{optbl}` ({n_op} champs métier) et la retourne."""
    db.define_table(
        "{optbl}",
{op_lines}
        *_meta_fields(),
        migrate=True,
    )
    return db.{optbl}


def define_allotissements(db):
    """Définit la table `{lottbl}` ({n_lot} champs métier, 1-N) et la retourne."""
    db.define_table(
        "{lottbl}",
        Field("operation", "reference {optbl}", label="Dossier opération", comment="Référence vers la table parente"),
        Field("objet_lot", "string", label="Objet du lot", comment="Objet du lot"),
        Field("numero_lot", "integer", label="N° du lot", comment="Numéro du lot"),
{lot_lines}
        *_meta_fields(),
        migrate=True,
    )
    return db.{lottbl}


def define_all(db):
    """Définit `{optbl}` puis sa sous-table `{lottbl}` (1-N)."""
    define_operations(db)
    define_allotissements(db)
    return db
'''


if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'align'
    mode = sys.argv[2] if len(sys.argv) > 2 else 'aor_pi'
    {'align': align, 'match': match, 'gen': gen}[cmd](mode)
